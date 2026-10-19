# SPDX-FileCopyrightText: 2026 Benoît Dupont de Dinechin
# SPDX-License-Identifier: Apache-2.0
"""svg2pptx — draw a project diagram SVG onto a python-pptx slide as native shapes.

Not a general SVG renderer. It covers exactly the subset the docs/diagrams/D*.svg files use,
which is deliberately narrow so that the result is editable in PowerPoint: every rect becomes a
rounded/square autoshape, every line a straight connector, every polygon a freeform, every text a
text box. Nothing is rasterised and nothing is embedded as a picture.

Supported: <svg viewBox>, <g> attribute inheritance, <rect>, <line>, <polygon>, <path> (M/L only),
<text>; presentation attributes fill, stroke, stroke-width, stroke-dasharray, font-family,
font-size, font-weight, text-anchor. Transforms, gradients, tspan and CSS are NOT supported — the
diagrams avoid them on purpose (they also have to survive PowerPoint's own SVG import).
"""

import re
import xml.etree.ElementTree as ET

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Pt

SVG_NS = "{http://www.w3.org/2000/svg}"
EMU_PER_PT = 12700

# Attributes a <g> passes down to its children.
INHERITED = (
    "fill", "stroke", "stroke-width", "stroke-dasharray",
    "font-family", "font-size", "font-weight", "text-anchor",
)

NAMED = {"black": "#000000", "white": "#ffffff", "none": None}


def _colour(value):
    """'#rrggbb' | '#rgb' | name -> RGBColor, or None for none/absent."""
    if not value:
        return None
    v = value.strip().lower()
    if v in NAMED:
        v = NAMED[v]
        if v is None:
            return None
    if not v.startswith("#"):
        return None
    h = v[1:]
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    if len(h) != 6:
        return None
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _num(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _dash_style(dasharray):
    """Map an SVG stroke-dasharray onto the nearest PowerPoint preset."""
    if not dasharray:
        return None
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    parts = [_num(p) for p in re.split(r"[,\s]+", dasharray.strip()) if p]
    if not parts:
        return None
    on = parts[0]
    return MSO_LINE_DASH_STYLE.ROUND_DOT if on <= 2 else MSO_LINE_DASH_STYLE.DASH


class Renderer:
    """Maps SVG user units onto slide EMU and emits shapes."""

    def __init__(self, shapes, scale, ox, oy):
        self.shapes = shapes
        self.s = scale          # EMU per SVG user unit
        self.ox, self.oy = ox, oy

    # -- unit helpers ------------------------------------------------------
    def X(self, v):
        return Emu(int(round(self.ox + v * self.s)))

    def Y(self, v):
        return Emu(int(round(self.oy + v * self.s)))

    def L(self, v):
        return Emu(max(1, int(round(v * self.s))))

    def pt(self, v):
        return Pt(round(v * self.s / EMU_PER_PT, 2))

    # -- styling -----------------------------------------------------------
    def _apply_fill(self, shape, style):
        rgb = _colour(style.get("fill"))
        if rgb is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb

    def _apply_line(self, obj, style, default_width=1.0):
        rgb = _colour(style.get("stroke"))
        if rgb is None:
            obj.line.fill.background()
            return
        obj.line.color.rgb = rgb
        obj.line.width = self.L(_num(style.get("stroke-width"), default_width))
        dash = _dash_style(style.get("stroke-dasharray"))
        if dash is not None:
            obj.line.dash_style = dash

    # -- element handlers --------------------------------------------------
    def rect(self, el, style):
        x, y = _num(el.get("x")), _num(el.get("y"))
        w, h = _num(el.get("width")), _num(el.get("height"))
        if w <= 0 or h <= 0:
            return
        rx = _num(el.get("rx"))
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rx > 0 else MSO_SHAPE.RECTANGLE
        shape = self.shapes.add_shape(kind, self.X(x), self.Y(y), self.L(w), self.L(h))
        shape.shadow.inherit = False
        if rx > 0:
            shape.adjustments[0] = min(0.5, rx / min(w, h))
        self._apply_fill(shape, style)
        self._apply_line(shape, style)
        shape.text_frame.word_wrap = False

    def line(self, el, style, pts=None):
        if pts is None:
            pts = [(_num(el.get("x1")), _num(el.get("y1"))),
                   (_num(el.get("x2")), _num(el.get("y2")))]
        for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
            c = self.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, self.X(x1), self.Y(y1), self.X(x2), self.Y(y2))
            c.shadow.inherit = False
            self._apply_line(c, style)

    def polygon(self, el, style):
        raw = [_num(p) for p in re.split(r"[,\s]+", (el.get("points") or "").strip()) if p]
        pts = list(zip(raw[0::2], raw[1::2]))
        if len(pts) < 3:
            return
        builder = self.shapes.build_freeform(self.X(pts[0][0]), self.Y(pts[0][1]))
        builder.add_line_segments([(self.X(x), self.Y(y)) for x, y in pts[1:]], close=True)
        shape = builder.convert_to_shape()
        shape.shadow.inherit = False
        self._apply_fill(shape, style)
        self._apply_line(shape, style)

    def path(self, el, style):
        """Open polyline: only M/L commands appear in these diagrams."""
        d = el.get("d") or ""
        if re.search(r"[^MLmlZz0-9.,\s+-]", d):
            raise ValueError(f"unsupported path command in {d!r}")
        tokens = re.findall(r"([MLml])|(-?[\d.]+)", d)
        pts, nums, cmd = [], [], None
        for letter, number in tokens:
            if letter:
                cmd = letter
                nums = []
            else:
                nums.append(float(number))
                if len(nums) == 2:
                    x, y = nums
                    if cmd in "ml" and pts:
                        x, y = pts[-1][0] + x, pts[-1][1] + y
                    pts.append((x, y))
                    nums = []
                    if cmd == "M":
                        cmd = "L"
                    elif cmd == "m":
                        cmd = "l"
        if len(pts) >= 2:
            self.line(el, style, pts=pts)

    def text(self, el, style):
        content = "".join(el.itertext()).strip()
        if not content:
            return
        fs = _num(style.get("font-size"), 12)
        x, y = _num(el.get("x")), _num(el.get("y"))
        anchor = style.get("text-anchor", "start")

        # Width is only an estimate; word_wrap is off, so text that overflows still renders
        # correctly. The box is sized so it is convenient to grab in PowerPoint.
        w = max(len(content) * fs * 0.62 + 8, 4 * fs)
        h = fs * 1.45
        left = x - w / 2 if anchor == "middle" else x
        # SVG y is the baseline; cap-height text sits visually ~0.36em above it.
        top = (y - 0.36 * fs) - h / 2

        box = self.shapes.add_textbox(self.X(left), self.Y(top), self.L(w), self.L(h))
        tf = box.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER if anchor == "middle" else PP_ALIGN.LEFT
        run = para.add_run()
        run.text = content
        font = run.font
        font.size = self.pt(fs)
        font.name = (style.get("font-family") or "Arial").split(",")[0].strip()
        font.bold = style.get("font-weight") in ("bold", "700", "800", "900")
        rgb = _colour(style.get("fill")) or RGBColor(0x22, 0x22, 0x22)
        font.color.rgb = rgb

    # -- traversal ---------------------------------------------------------
    def walk(self, node, inherited):
        for el in node:
            tag = el.tag.replace(SVG_NS, "")
            style = dict(inherited)
            style.update({k: v for k, v in el.attrib.items() if k in INHERITED})
            if tag == "g":
                self.walk(el, style)
            elif tag == "rect":
                self.rect(el, style)
            elif tag == "line":
                self.line(el, style)
            elif tag == "polygon":
                self.polygon(el, style)
            elif tag == "path":
                self.path(el, style)
            elif tag == "text":
                self.text(el, style)
            elif tag in ("title", "desc", "defs", "metadata", "style"):
                continue
            else:
                raise ValueError(f"unsupported SVG element <{tag}>")


def draw(svg_path, slide, box, drop_background=True):
    """Draw svg_path onto slide, fitted inside box = (left, top, width, height) in EMU.

    The SVG's full-canvas white backing rect is dropped by default: the slide already has a
    background, and an opaque rect over the whole frame is a nuisance to click past in PowerPoint.
    Returns the number of shapes added.
    """
    root = ET.parse(svg_path).getroot()
    vb = (root.get("viewBox") or "").split()
    if len(vb) == 4:
        _, _, vw, vh = (float(v) for v in vb)
    else:
        vw, vh = _num(root.get("width"), 1000), _num(root.get("height"), 600)

    left, top, width, height = box
    scale = min(width / vw, height / vh)
    ox = left + (width - vw * scale) / 2
    oy = top + (height - vh * scale) / 2

    if drop_background:
        for el in list(root):
            if el.tag.replace(SVG_NS, "") == "rect" and _num(el.get("width")) >= vw:
                root.remove(el)
                break

    before = len(slide.shapes._spTree)
    Renderer(slide.shapes, scale, ox, oy).walk(root, {})
    return len(slide.shapes._spTree) - before
